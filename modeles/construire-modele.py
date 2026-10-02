# -*- coding: utf-8 -*-
"""Fabrique « Mon-Dje-modele-produits.xlsx » : le fichier Excel que les commerces remplissent pour importer leur liste
de produits (et leur stock de départ) dans Mon Djê.

    python cms/modeles/construire-modele.py

Les en-têtes de la feuille « Mes produits » sont EXACTEMENT ceux que lit maquis/src/importer.ts (MODELE_IMPORT) :
Nom, Unité, Prix de vente, Prix d’achat, Stock. Un export CSV ne garde que la feuille active : « Mes produits » doit
rester la première. Elle est vide (seulement les en-têtes) pour qu'aucun exemple ne soit importé par mégarde.
"""
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

ICI = os.path.dirname(os.path.abspath(__file__))
SORTIE = os.path.join(ICI, "Mon-Dje-modele-produits.xlsx")

NUIT, ACCENT, CIEL, GRIS = "0E1B2E", "F4A300", "E6F4FE", "5B667A"
POLICE = "Calibri"
LIGNES = 500          # = MAX_LIGNES de l'importeur

ENTETES = ["Nom", "Unité", "Prix de vente", "Prix d’achat", "Stock"]
OBLIGATOIRES = {"Nom", "Prix de vente"}
UNITES = ["pièce", "bouteille", "canette", "casier", "carton", "sac", "sachet", "paquet", "kg", "litre",
          "assiette", "dose", "tas", "détail", "demi", "gros", "article"]
AIDE = {
    "Nom": "Obligatoire. 2 à 80 caractères. Un nom par ligne, sans doublon (« Castel Beer 33 cl »).",
    "Unité": "Facultatif. Choisissez dans la liste ou écrivez la vôtre (bouteille, sac, kg…). Si vide : « pièce ».",
    "Prix de vente": "Obligatoire. En francs CFA, nombre entier supérieur à 0, sans « F » (ex. 700).",
    "Prix d’achat": "Facultatif. Ce que vous payez au fournisseur, pour calculer votre marge (ex. 450).",
    "Stock": "Facultatif. La quantité que vous avez aujourd’hui, en nombre entier. Si vide : 0.",
}

wb = Workbook()
wb.properties.title = "Mon Djê — modèle pour importer mes produits"
wb.properties.creator = "Mon Djê"

# ---------------------------------------------------------------- Feuille 1 : à remplir
ws = wb.active
ws.title = "Mes produits"
fin = Side(style="thin", color="D5DBE5")
for c, nom in enumerate(ENTETES, start=1):
    cell = ws.cell(row=1, column=c, value=nom)
    obligatoire = nom in OBLIGATOIRES
    cell.font = Font(name=POLICE, bold=True, size=12, color=NUIT if obligatoire else "FFFFFF")
    cell.fill = PatternFill("solid", fgColor=ACCENT if obligatoire else NUIT)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="medium", color=NUIT))
    cell.comment = Comment(AIDE[nom], "Mon Djê", width=300, height=90)
ws.row_dimensions[1].height = 30
for lettre, largeur in zip("ABCDE", (44, 16, 16, 16, 12)):
    ws.column_dimensions[lettre].width = largeur
ws.freeze_panes = "A2"
for r in range(2, LIGNES + 2):
    for c in (3, 4, 5):
        ws.cell(row=r, column=c).number_format = "0"      # nombres entiers, sans séparateur de milliers
    ws.cell(row=r, column=1).font = Font(name=POLICE, size=11)
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = Border(bottom=fin)

zone = lambda col: f"{col}2:{col}{LIGNES + 1}"
v_nom = DataValidation(type="textLength", operator="between", formula1="2", formula2="80", allow_blank=True,
                       errorTitle="Nom", error="Le nom doit faire entre 2 et 80 caractères.", errorStyle="warning")
v_unite = DataValidation(type="list", formula1='"' + ",".join(UNITES) + '"', allow_blank=True, showErrorMessage=False)
v_pv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                      errorTitle="Prix de vente", error="Écrivez un nombre entier supérieur à 0, sans « F ».", errorStyle="warning")
v_pa = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                      errorTitle="Prix d’achat", error="Écrivez un nombre entier (0 ou plus), sans « F ».", errorStyle="warning")
v_stock = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         errorTitle="Stock", error="Écrivez un nombre entier (0 ou plus).", errorStyle="warning")
for v, col in ((v_nom, "A"), (v_unite, "B"), (v_pv, "C"), (v_pa, "D"), (v_stock, "E")):
    ws.add_data_validation(v)
    v.add(zone(col))

# ---------------------------------------------------------------- Feuille 2 : mode d'emploi
em = wb.create_sheet("Mode d'emploi")
em.sheet_view.showGridLines = False
em.column_dimensions["A"].width = 4
em.column_dimensions["B"].width = 36
em.column_dimensions["C"].width = 72
ligne = [1]


def ecrire(col, texte, **o):
    c = em.cell(row=ligne[0], column=col, value=texte)
    c.font = Font(name=POLICE, size=o.get("taille", 11), bold=o.get("gras", False), color=o.get("couleur", "1C2433"))
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if o.get("fond"):
        c.fill = PatternFill("solid", fgColor=o["fond"])
    return c


def titre(texte):
    ligne[0] += 1
    ecrire(2, texte, taille=14, gras=True, couleur=NUIT)
    em.merge_cells(start_row=ligne[0], start_column=2, end_row=ligne[0], end_column=3)
    em.row_dimensions[ligne[0]].height = 24
    ligne[0] += 1


def paire(a, b, hauteur=None, fond=None):
    ecrire(2, a, gras=True, fond=fond)
    ecrire(3, b, fond=fond)
    em.row_dimensions[ligne[0]].height = max(hauteur or 0, 32)
    ligne[0] += 1


ecrire(2, "Importer mes produits dans Mon Djê", taille=18, gras=True, couleur=NUIT)
em.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)
em.row_dimensions[1].height = 30
ligne[0] = 2
ecrire(2, "Remplissez la feuille « Mes produits » avec votre liste et votre stock de départ, puis importez-la dans l’application. "
          "Plus besoin de tout ressaisir au téléphone.", couleur=GRIS)
em.merge_cells(start_row=2, start_column=2, end_row=2, end_column=3)
em.row_dimensions[2].height = 34

titre("Les 4 étapes")
paire("1. Remplir", "Dans la feuille « Mes produits » : une ligne par produit. Les colonnes jaunes (Nom, Prix de vente) sont obligatoires, "
                    "les autres sont facultatives. Jusqu’à 500 lignes par fichier ; au-delà, faites un second fichier.", 62)
paire("2. Enregistrer en CSV", "Dans Excel : Fichier → Enregistrer sous → type « CSV UTF-8 (délimité par des virgules) » "
                               "(choisissez bien « UTF-8 », pour garder les accents). Dans Google Sheets : Fichier → Télécharger → Valeurs séparées par des virgules (.csv). "
                               "Excel prévient que seule la feuille active sera enregistrée : c’est normal, répondez « OK ».", 76)
paire("3. Envoyer sur le téléphone", "Envoyez le fichier .csv sur le téléphone qui a Mon Djê (WhatsApp, e-mail ou câble). "
                                     "Enregistrez-le dans les téléchargements.", 48)
paire("4. Importer", "Dans Mon Djê, connecté en propriétaire : Carte et stock → « Importer ma liste de produits (fichier Excel/CSV) » → "
                     "« Choisir mon fichier ». L’application vous montre ce qui passe et ce qui est refusé, avec la raison, "
                     "et n’écrit rien tant que vous n’avez pas touché « Importer ».", 76)

titre("Les colonnes")
paire("Nom *", "Obligatoire. 2 à 80 caractères. Deux noms qui ne diffèrent que par les majuscules, les accents ou les espaces sont le même produit.", 34, CIEL)
paire("Unité", "Facultatif. Une unité de la liste ou la vôtre. Si vide : « pièce ».", 20)
paire("Prix de vente *", "Obligatoire. En francs CFA, nombre entier supérieur à 0, sans le signe « F ». Exemple : 700.", 34, CIEL)
paire("Prix d’achat", "Facultatif. Le prix que vous payez au fournisseur : il sert à calculer votre marge. Exemple : 450.", 34)
paire("Stock", "Facultatif. La quantité que vous avez aujourd’hui. Elle est enregistrée comme une entrée de stock datée d’aujourd’hui. Si vide : 0.", 34)

titre("Si une ligne est refusée")
paire("nom manquant ou trop court", "La colonne Nom est vide ou fait moins de 2 caractères.", 20)
paire("prix de vente manquant ou invalide", "Le prix de vente est vide, à 0, ou n’est pas un nombre.", 20)
paire("ce produit existe déjà", "Le produit est déjà dans votre carte, ou en double dans le fichier. Un produit existant n’est jamais écrasé.", 34)
paire("accents abîmés", "Si « Bière » apparaît « Bi?re », enregistrez de nouveau en « CSV UTF-8 ».", 20)

titre("Bon à savoir")
for texte in (
    "• Ne changez pas les titres des colonnes de la feuille « Mes produits », et gardez-la en première position.",
    "• Les exemples de la feuille « Exemples » ne sont pas importés : copiez-les dans « Mes produits » si vous voulez vous en servir.",
    "• Vous pouvez importer plusieurs fois : les produits déjà présents sont simplement refusés.",
    "• Une question ? WhatsApp 05 65 75 03 03, ou mondje.ci/modeles.",
):
    ecrire(2, texte)
    em.merge_cells(start_row=ligne[0], start_column=2, end_row=ligne[0], end_column=3)
    em.row_dimensions[ligne[0]].height = 20
    ligne[0] += 1

# ---------------------------------------------------------------- Feuille 3 : exemples
ex = wb.create_sheet("Exemples")
ex.column_dimensions["A"].width = 38
for l, w in zip("BCDE", (14, 16, 16, 12)):
    ex.column_dimensions[l].width = w
ex["A1"] = "Exemples (non importés) — copiez-les dans « Mes produits » si besoin"
ex["A1"].font = Font(name=POLICE, bold=True, size=12, color=NUIT)
ex.merge_cells("A1:E1")
BLOCS = [
    ("Maquis / restaurant", [
        ("Castel Beer 33 cl", "bouteille", 700, 450, 24), ("Coca-Cola 33 cl", "bouteille", 500, 330, 48),
        ("Poulet braisé", "assiette", 3500, 2200, 0), ("Attiéké poisson", "assiette", 2500, 1500, 0)]),
    ("Boutique", [
        ("Riz parfumé 5 kg", "sac", 4500, 3800, 10), ("Huile 1 L", "litre", 1300, 1050, 24),
        ("Sucre en poudre", "kg", 800, 650, 30), ("Savon de ménage", "pièce", 300, 220, 60)]),
    ("Marché", [
        ("Tomates", "tas", 500, 300, 20), ("Oignons", "tas", 500, 320, 15),
        ("Attiéké", "sachet", 250, 150, 40), ("Piment frais", "kg", 1500, 1000, 5)]),
]
r = 3
for nom, produits in BLOCS:
    ex.cell(row=r, column=1, value=nom).font = Font(name=POLICE, bold=True, color=NUIT)
    ex.cell(row=r, column=1).fill = PatternFill("solid", fgColor=CIEL)
    for c in range(2, 6):
        ex.cell(row=r, column=c).fill = PatternFill("solid", fgColor=CIEL)
    r += 1
    for c, titre_col in enumerate(ENTETES, start=1):
        cell = ex.cell(row=r, column=c, value=titre_col)
        cell.font = Font(name=POLICE, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NUIT)
    r += 1
    for p in produits:
        for c, val in enumerate(p, start=1):
            ex.cell(row=r, column=c, value=val).font = Font(name=POLICE)
        r += 1
    r += 1

# Impression : une page de large pour chaque feuille
for feuille in wb.worksheets:
    feuille.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    feuille.page_setup.fitToWidth = 1
    feuille.page_setup.fitToHeight = 0
    feuille.page_setup.paperSize = 9   # A4

wb.active = 0
wb.save(SORTIE)
print("écrit", SORTIE)
