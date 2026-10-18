# -*- coding: utf-8 -*-
"""Fabrique les fichiers Excel PAR MÉTIER : « Mon-Dje-modele-produits-<métier>.xlsx » (lot B, 18/10/2026).

Chaque fichier est le modèle général (cms/modeles/Mon-Dje-modele-produits.xlsx : feuille « Mes produits » vide, mode d'emploi, exemples) auquel s'ajoute une feuille
« Liste du métier » : tous les produits du catalogue Mon Djê pour ce métier (nom, unité, rubrique), prix à remplir. Le commerçant efface les lignes qu'il ne vend pas, écrit ses prix
et son stock, puis copie A:E dans « Mes produits » (l'importation ne lit que cette feuille : les lignes sans prix y seraient refusées, d'où cette feuille à part).

    python -X utf8 cms/modeles/construire-modeles-metiers.py <catalogue.json>

<catalogue.json> : le catalogue actuel, une ligne par (activité, produit) : activite, rubrique, rang, nom, contenance, unite, conditionnement, taille (voir outils/catalogue/lire_catalogue.sql).
"""
import json, os, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.properties import PageSetupProperties

ICI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(ICI, 'Mon-Dje-modele-produits.xlsx')
NUIT, ACCENT, CIEL, GRIS, POLICE = '0E1B2E', 'F4A300', 'E6F4FE', '5B667A', 'Calibri'

METIERS = {   # activité du catalogue -> (nom du fichier, titre)
    'boulangerie': ('boulangerie', 'Boulangerie, pâtisserie'), 'marche': ('marche', 'Marché'), 'boucherie': ('boucherie-poissonnerie', 'Boucherie, poissonnerie'),
    'depot_boissons': ('depot-de-boissons', 'Dépôt de boissons'), 'quincaillerie': ('quincaillerie', 'Quincaillerie, matériaux'),
    'cosmetiques': ('cosmetiques-bijouterie', 'Cosmétiques, bijouterie'), 'papeterie': ('papeterie-librairie', 'Papeterie, librairie'),
    'telephonie': ('telephones-accessoires', 'Téléphones et accessoires'), 'vetements': ('vetements-friperie-chaussures', 'Vêtements, friperie, chaussures'),
    'pieces_auto': ('pieces-auto-moto', 'Pièces auto et moto'), 'boutique': ('boutique', 'Boutique'), 'bar_maquis': ('bar-maquis', 'Bar / Maquis'),
    'restaurant': ('restaurant', 'Restaurant'),
}

lignes = json.load(open(sys.argv[1], encoding='utf-8'))
os.makedirs(os.path.join(ICI, 'metiers'), exist_ok=True)
for activite, (fichier, titre) in METIERS.items():
    # Les rubriques propres au métier d'abord (pour la boulangerie : celles de la famille « Boulangerie et pâtisserie », rang 8xxxx), puis le reste dans l'ordre du catalogue.
    produits = sorted((l for l in lignes if l['activite'] == activite), key=lambda l: (0 if (activite == 'boulangerie' and l['rang'] >= 80000) else 1, l['rang'], l['nom'], l['contenance'] or ''))
    wb = openpyxl.load_workbook(BASE)
    ws = wb.create_sheet('Liste du métier', 1)
    ws['A1'] = f'Liste du métier : {titre}'
    ws['A1'].font = Font(name=POLICE, bold=True, size=14, color=NUIT)
    ws.merge_cells('A1:F1')
    ws['A2'] = ('Voici les produits que Mon Djê connaît pour ce métier. 1) Effacez les lignes que vous ne vendez pas. 2) Écrivez votre PRIX DE VENTE (obligatoire), et si vous voulez le prix d’achat et le stock. '
                '3) Copiez les colonnes A à E (de « Nom » à « Stock ») et collez-les dans la feuille « Mes produits », à partir de la ligne 2. 4) Enregistrez en CSV et importez (voir « Mode d’emploi »). '
                'Il n’y a AUCUN prix ici : ils changent d’un quartier à l’autre. Un produit qui manque ? Ajoutez-le vous-même au bas de la feuille.')
    ws['A2'].font = Font(name=POLICE, size=11, color=GRIS)
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells('A2:F2')
    ws.row_dimensions[2].height = 92
    entetes = ['Nom', 'Unité', 'Prix de vente', 'Prix d’achat', 'Stock', 'Rubrique (pour vous repérer, ne pas copier)']
    for c, nom in enumerate(entetes, start=1):
        cell = ws.cell(row=4, column=c, value=nom)
        obligatoire = nom in ('Nom', 'Prix de vente')
        cell.font = Font(name=POLICE, bold=True, size=12, color=NUIT if obligatoire else 'FFFFFF')
        cell.fill = PatternFill('solid', fgColor=ACCENT if obligatoire else NUIT)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[4].height = 30
    r, derniere = 5, None
    for p in produits:
        nom = f"{p['nom']} {p['contenance']}" if p['contenance'] else p['nom']
        if p['rubrique'] != derniere:
            derniere = p['rubrique']
        ws.cell(row=r, column=1, value=nom).font = Font(name=POLICE)
        ws.cell(row=r, column=2, value=p['unite']).font = Font(name=POLICE)
        for c in (3, 4, 5):
            ws.cell(row=r, column=c).number_format = '0'
        ws.cell(row=r, column=6, value=p['rubrique']).font = Font(name=POLICE, color=GRIS, italic=True)
        r += 1
    for l, w in zip('ABCDEF', (50, 14, 16, 16, 12, 46)):
        ws.column_dimensions[l].width = w
    ws.freeze_panes = 'A5'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.page_setup.paperSize = 9
    wb.properties.title = f'Mon Djê — modèle pour importer mes produits : {titre}'
    sortie = os.path.join(ICI, 'metiers', f'Mon-Dje-modele-produits-{fichier}.xlsx')
    wb.save(sortie)
    print(f'{titre} : {len(produits)} produits -> {os.path.basename(sortie)}')
